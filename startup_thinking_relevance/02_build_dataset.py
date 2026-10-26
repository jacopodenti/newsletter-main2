"""
STEP 2 — Combina le fonti etichettate in un dataset unico, con testo intero.

Composizione:
  positivi (1): 161 "editorial_startup" da SUT_Elenco news.xlsx (25 fogli settimanali,
    tutti trattati come pertinenti — sono già stati pubblicati in newsletter passate)
  negativi/positivi: 100 "llm_judged" da negative_candidates.xlsx, etichetta presa dalla
    colonna "Giudizio LLM" (NESSUNA revisione umana fatta su questi — tracciato in origin)

Ogni riga traccia l'ORIGINE. Scarica il testo intero dagli URL (fallback al sommario).

Output: data/dataset.jsonl + data/dataset_meta.json
Uso:    python startup_thinking_relevance/02_build_dataset.py
"""
import json
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "ocx_relevance"))
from common import fetch_fulltext, log

DATA = HERE / "data"
XLSX_EDITORIAL = HERE / "SUT_Elenco news.xlsx"
XLSX_NEGATIVES = DATA / "negative_candidates.xlsx"


def load_editorial():
    # SUT_Elenco news.xlsx: un foglio per settimana, colonne A-F =
    # Data, Fonte, LINK, Titolo, Riassunto, CATEGORIA — header in riga 1
    wb = openpyxl.load_workbook(XLSX_EDITORIAL, data_only=True)
    out, seen = [], set()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for r in range(2, ws.max_row + 1):
            title = ws.cell(row=r, column=4).value
            if not title or not str(title).strip():
                continue
            link = str(ws.cell(row=r, column=3).value or "").strip()
            if link and link in seen:
                continue
            seen.add(link)
            out.append({
                "title": str(title).strip(),
                "summary": str(ws.cell(row=r, column=5).value or ""),
                "link": link,
                "label": 1,
                "origin": "editorial_startup",
            })
    return out


def load_negatives():
    wb = openpyxl.load_workbook(XLSX_NEGATIVES, data_only=True)
    ws = wb["Da giudicare"]
    out = []
    for r in range(2, ws.max_row + 1):
        title = ws.cell(row=r, column=1).value
        if not title:
            continue
        giudizio = str(ws.cell(row=r, column=6).value or "").strip().lower()
        if giudizio not in ("si", "no"):
            continue  # salta eventuali righe ERRORE
        out.append({
            "title": str(title).strip(),
            "summary": str(ws.cell(row=r, column=5).value or ""),
            "link": str(ws.cell(row=r, column=4).value or "").strip(),
            "label": 1 if giudizio == "si" else 0,
            "origin": "llm_judged",  # nessuna revisione umana — tracciato esplicitamente
        })
    return out


def main():
    log("STEP 2 — costruzione dataset startup-thinking (editorial + llm_judged)")
    # 1) carica le due fonti etichettate e le combina, tracciando l'ORIGINE di ogni riga
    ed = load_editorial()
    neg = load_negatives()
    rows = ed + neg

    log(f"  editorial_startup: {len(ed):>4}  (tutti positivi)")
    llm_si = sum(x["label"] for x in neg)
    log(f"  llm_judged:   {len(neg):>4}  (Si={llm_si}, No={len(neg)-llm_si} — etichetta = giudizio LLM, non revisionata)")
    log(f"  TOTALE:       {len(rows):>4}  (positivi={sum(x['label'] for x in rows)}, negativi={sum(1-x['label'] for x in rows)})")

    # 2) scarica il testo intero dagli URL (fallback al sommario se il fetch fallisce)
    log("\n  scarico il testo intero dagli URL...")
    def fetch(x):
        text, src = fetch_fulltext(x["link"], x["summary"])
        x["full_text"] = text
        x["text_source"] = src
        return x
    with ThreadPoolExecutor(max_workers=16) as ex:
        rows = list(ex.map(fetch, rows))
    nfull = sum(1 for x in rows if x["text_source"] == "full")
    log(f"  testo intero: {nfull}/{len(rows)} ({nfull/len(rows):.0%})  fallback sommario: {len(rows)-nfull}")

    # 3) output — dataset.jsonl (righe) + dataset_meta.json (statistiche riassuntive)
    with open(DATA / "dataset.jsonl", "w", encoding="utf-8") as f:
        for i, x in enumerate(rows):
            x["id"] = i
            f.write(json.dumps(x, ensure_ascii=False) + "\n")

    meta = {
        "totale": len(rows),
        "positivi": sum(x["label"] for x in rows),
        "negativi": sum(1 - x["label"] for x in rows),
        "per_origine": {o: sum(1 for x in rows if x["origin"] == o) for o in {x["origin"] for x in rows}},
        "testo_intero": nfull, "fallback_sommario": len(rows) - nfull,
    }
    json.dump(meta, open(DATA / "dataset_meta.json", "w"), ensure_ascii=False, indent=2)
    log(f"\n  salvato: {DATA/'dataset.jsonl'}  ({len(rows)} righe)")
    log(f"  salvato: {DATA/'dataset_meta.json'}")
    log("\nFATTO.")


if __name__ == "__main__":
    main()
