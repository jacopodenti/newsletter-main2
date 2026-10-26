"""
STEP 1 — Prepara un campione di candidati "probabilmente negativi" da far
giudicare a mano, con un giudizio LLM già proposto come punto di partenza.

Selezione: campiona dal corpus grezzo condiviso escludendo i feed dedicati a
startup/open innovation (altrimenti il campione sarebbe troppo pieno di veri
positivi), con un tetto per fonte per evitare che una sola testata domini il
campione (lo stesso problema di concentrazione trovato nel test OCX).

Poi ogni candidato viene fatto giudicare a un LLM (Si/No + motivo breve),
usando la descrizione e i criteri di scope dell'Osservatorio Startup Thinking
(da "Descrizione Oss. Startup Thinking[44].rtf") — il giudizio LLM è un PUNTO
DI PARTENZA da confermare o correggere a mano, non una decisione finale.

Output: data/negative_candidates.xlsx (100 righe, pronte per la revisione umana)
Uso:    python startup_thinking_relevance/01_build_negative_candidates.py
"""
import json
import random
import re
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from dotenv import load_dotenv
import os

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
load_dotenv(ROOT / ".env")

DATA = HERE / "data"
N_CANDIDATES = 100
MAX_PER_FEED = 3
SEED = 42

# feed troppo dedicati a startup/open innovation: esclusi dal campionamento
# "probabilmente negativo" (altrimenti il campione sarebbe pieno di veri positivi)
STARTUP_HEAVY_FEEDS = {
    "StartupItalia", "I3P", "EU-Startups", "Crunchbase News", "Sifted",
    "Tech.eu", "The Next Web", "Global Corporate Venturing", "The Innovator",
}

STARTUP_DESCRIPTION = """L'Osservatorio Startup Thinking si occupa di diffondere nelle imprese l'Open Innovation e la cultura imprenditoriale attraverso la contaminazione con l'ecosistema startup e con la community dell'innovazione. La ricerca analizza come le imprese consolidate possano generare nuove fonti di valore attraverso un Innovation Management strutturato, aperto e imprenditoriale.

Due direttrici: (1) Spinta all'Open Innovation — coinvolgimento di startup, scaleup, università, incubatori, acceleratori, investitori, partner industriali attraverso scouting, partnership, sperimentazioni, investimenti; (2) Spinta alla trasformazione interna — capacità dell'impresa di organizzare, governare ed eseguire l'innovazione (strategia, governance, cultura, competenze, processi, portafoglio iniziative, metriche)."""

PROMPT = """Sei un classificatore di pertinenza per l'Osservatorio "Startup Thinking" del Politecnico di Milano.

DESCRIZIONE OSSERVATORIO:
{desc}

COMPITO: dato un articolo, decidi se e PERTINENTE a questo osservatorio.
RIENTRA NELLO SCOPE quando il tema centrale è: una corporate che avvia/modifica una strategia o funzione di innovazione; collaborazione corporate-startup/ecosistema; programmi di scouting, Open Innovation, venture clienting, accelerazione, incubazione; PoC/pilot/progetti sperimentali; criteri o metriche per selezionare/scalare iniziative; cambiamenti in ruoli, governance, cultura dell'innovazione; Corporate Venture Capital, corporate venturing, venture building, spin-off, acquisizioni di startup con finalità strategica; ecosistema startup, investimenti, crescita scaleup, trasferimento tecnologico, policy di sistema; una tecnologia emergente analizzata tramite casi d'uso/startup/adozione aziendale/impatti organizzativi.
NON RIENTRA (o rilevanza molto bassa): "innovativo" usato solo come aggettivo promozionale; lanci prodotto/recensioni senza informazioni sul processo di innovazione o adozione aziendale; annunci di modelli AI/benchmark/funzionalità senza casi d'uso aziendali o implicazioni per l'Innovation Management; articoli tecnici specialistici senza collegamento a startup/imprese/sperimentazione; cronaca finanziaria ordinaria (titoli, trimestrali, dividendi) senza componente di innovazione; M&A motivate solo da consolidamento/quota di mercato; notizie di marketing/CX/sostenibilità/cybersecurity/HR dove l'innovazione organizzativa non è centrale; nomine manageriali non legate all'innovazione; eventi/premi/classifiche senza dati o iniziative concrete; singoli round di investimento di rilevanza limitata senza informazioni su trend/ecosistema/scalabilità; contenuti motivazionali/teorici sull'imprenditorialità senza casi concreti.

Rispondi SOLO con JSON: {{"rilevante": true/false, "confidence": 0.0-1.0, "motivo": "una frase breve in italiano"}}

ARTICOLO:
Titolo: {title}
Sommario: {summary}

JSON:"""


def strip_html(text: str) -> str:
    return re.sub(r"<[^>]+>", " ", text or "").strip()


def build_client():
    from openai import OpenAI
    return OpenAI(
        api_key=os.environ["AZURE_OPENAI_API_KEY"],
        base_url=os.environ["AZURE_OPENAI_ENDPOINT"],
    )


def judge(client, deployment, article):
    summary = strip_html(article.get("summary", ""))[:1200]
    prompt = PROMPT.format(desc=STARTUP_DESCRIPTION, title=article["title"], summary=summary)
    for attempt in range(5):
        try:
            resp = client.responses.create(model=deployment, input=prompt)
            txt = resp.output_text.strip()
            if txt.startswith("```"):
                txt = txt.split("```")[1].removeprefix("json").strip()
            res = json.loads(txt)
            return {
                "rilevante": bool(res.get("rilevante")),
                "confidence": float(res.get("confidence", 0.0)),
                "motivo": res.get("motivo", ""),
            }
        except Exception:
            if attempt == 4:
                return {"rilevante": None, "confidence": 0.0, "motivo": "ERRORE"}
            time.sleep(2 ** attempt)


def sample_candidates(corpus_path, n, max_per_feed, seed):
    rows = [json.loads(l) for l in open(corpus_path, encoding="utf-8")]
    rows = [r for r in rows if r["feed_name"] not in STARTUP_HEAVY_FEEDS and r["title"]]
    random.Random(seed).shuffle(rows)

    per_feed = {}
    picked = []
    for r in rows:
        f = r["feed_name"]
        if per_feed.get(f, 0) >= max_per_feed:
            continue
        per_feed[f] = per_feed.get(f, 0) + 1
        picked.append(r)
        if len(picked) >= n:
            break
    return picked


def main():
    # 1) campiona candidati "probabilmente negativi" dal corpus grezzo (esclude i feed
    # startup-heavy, tetto per fonte per evitare che una sola testata domini il campione)
    corpus_path = ROOT / "corpus" / "corpus_raw.jsonl"
    print(f"STEP 1 — campiono {N_CANDIDATES} candidati 'probabilmente negativi' (max {MAX_PER_FEED}/fonte)")
    candidates = sample_candidates(corpus_path, N_CANDIDATES, MAX_PER_FEED, SEED)
    print(f"  campionati: {len(candidates)} da {len(set(c['feed_name'] for c in candidates))} fonti diverse")

    # 2) giudizio LLM (Si/No + motivo) su ogni candidato, in parallelo — punto di partenza
    # da confermare/correggere a mano, non una decisione finale
    deployment = os.environ.get("AZURE_OPENAI_DEPLOYMENT", "gpt-5.4-nano")
    client = build_client()

    print(f"  giudizio LLM ({deployment}) su ognuno...")
    with ThreadPoolExecutor(max_workers=8) as ex:
        judged = list(ex.map(lambda a: judge(client, deployment, a), candidates))

    si = sum(1 for j in judged if j["rilevante"] is True)
    no = sum(1 for j in judged if j["rilevante"] is False)
    err = sum(1 for j in judged if j["rilevante"] is None)
    print(f"  LLM: Si={si}  No={no}  errori={err}")

    # 3) foglio Excel di output, pronto per la revisione umana (colonna "Giudizio umano"
    # lasciata vuota apposta)
    out_path = DATA / "negative_candidates.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Da giudicare"
    heads = ["Titolo", "Fonte", "Data", "Link", "Sommario", "Giudizio LLM", "Confidence LLM", "Motivo LLM", "Giudizio umano (Si/No)"]
    hf = Font(bold=True, color="FFFFFF")
    hd = PatternFill("solid", fgColor="1F6F3E")
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = hf
        cell.fill = hd

    for i, (a, j) in enumerate(zip(candidates, judged), 2):
        ws.cell(row=i, column=1, value=a["title"])
        ws.cell(row=i, column=2, value=a["feed_name"])
        ws.cell(row=i, column=3, value=a["published_date"])
        ws.cell(row=i, column=4, value=a["link"])
        ws.cell(row=i, column=5, value=strip_html(a.get("summary", ""))[:500])
        ws.cell(row=i, column=6, value="Si" if j["rilevante"] is True else ("No" if j["rilevante"] is False else "ERRORE"))
        ws.cell(row=i, column=7, value=j["confidence"])
        ws.cell(row=i, column=8, value=j["motivo"])
        # colonna 9 lasciata vuota apposta: da compilare a mano

    widths = {"A": 55, "B": 22, "C": 14, "D": 40, "E": 55, "F": 12, "G": 14, "H": 45, "I": 22}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    wb.save(out_path)
    print(f"\n  salvato: {out_path}")
    print("\nFATTO.")


if __name__ == "__main__":
    main()
