"""
STEP 9 — Compositore: seleziona gli articoli per il numero, non riordina e non rivaluta la
pertinenza. Legge data/ranking.csv (già prodotto da 06_rank.py, con content_type/target_audience
già taggati sui candidati Si E No visti dal gate), applica i vincoli di
config/composer_ocx.json, esclude chi è già in store/articles.jsonl con
selected_for_newsletter=true, e scrive data/composizione.xlsx.

NON aggiorna lo store: la selezione qui prodotta è una PROPOSTA da rivedere, non un invio
confermato — l'aggiornamento di selected_for_newsletter resta un passo separato, a valle della
revisione umana.

Uso:
  python ocx_relevance/09_compose.py
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
import csv

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "ocx_relevance"))
from common import log, canonicalize_url, load_tracking_params, select_for_newsletter

DATA = HERE / "data"
OSSERVATORIO = "ocx"
CONFIG_PATH = ROOT / "config" / "composer_ocx.json"


def load_candidates():
    """Solo le righe di ranking.csv viste dal gate (pertinenza_llm non vuota) — sono le sole
    ad avere content_type/target_audience taggati. Il resto del corpus non è mai stato
    giudicato, quindi non è un candidato eleggibile per il compositore."""
    tp = load_tracking_params()
    rows = list(csv.DictReader(open(DATA / "ranking.csv", encoding="utf-8")))
    candidates = []
    for r in rows:
        if not r.get("pertinenza_llm"):
            continue
        candidates.append({
            "title": r["titolo"],
            "rank": int(r["rank"]),
            "pertinenza_llm": r["pertinenza_llm"],
            "gate_confidence": float(r["confidence_llm"]) if r.get("confidence_llm") else None,
            "fonte": r["fonte"],
            "content_type": r.get("content_type") or None,
            "target_audience": r.get("target_audience") or None,
            "link": r["link"],
            "url_canonico": canonicalize_url(r["link"], tp),
        })
    return candidates


def load_already_selected():
    store_path = ROOT / "store" / "articles.jsonl"
    if not store_path.exists():
        return set()
    rows = [json.loads(l) for l in open(store_path, encoding="utf-8")]
    return {r["url_canonico"] for r in rows
            if r.get("osservatorio") == OSSERVATORIO and r.get("selected_for_newsletter")}


def write_output(selected, esclusi, log_allentamento, config):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Selezionati"
    hd = Font(bold=True, color="FFFFFF")
    hd_fill = PatternFill("solid", fgColor="1F6F3E")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")

    note = f"Compositore OCX — {len(selected)}/{config['n_articoli_totali']} slot riempiti."
    if any("allentato" in l or "BLOCCATO" in l or "ripescaggio:" in l for l in log_allentamento):
        note += "  ATTENZIONE: " + " | ".join(log_allentamento)
    ws["A1"] = note
    ws["A1"].font = Font(bold=True, size=11)
    if "BLOCCATO" in " ".join(log_allentamento):
        ws["A1"].fill = warn_fill

    heads = ["Rank", "Via", "Content type", "Testata", "Gruppo editoriale", "Target audience (informativo)", "Titolo", "Link"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.fill = hd, hd_fill
    for row, a in enumerate(selected, start=4):
        ws.cell(row=row, column=1, value=a["rank"])
        ws.cell(row=row, column=2, value=a["_via"])
        ws.cell(row=row, column=3, value=a.get("content_type") or "non_classificato")
        ws.cell(row=row, column=4, value=a["fonte"])
        ws.cell(row=row, column=5, value=a["_gruppo"])
        ws.cell(row=row, column=6, value=a.get("target_audience") or "non_classificato")
        ws.cell(row=row, column=7, value=a["title"])
        c8 = ws.cell(row=row, column=8, value="apri articolo")
        c8.hyperlink = a["link"]
        c8.font = Font(color="0563C1", underline="single")

    ws2 = wb.create_sheet("Esclusi")
    ws2["A1"] = f"{len(esclusi)} articoli visti dal gate ma non selezionati — motivo in colonna E, per ripescare a mano."
    ws2["A1"].font = Font(bold=True, size=11)
    heads2 = ["Rank", "Content type", "Testata", "Gruppo editoriale", "Motivo esclusione", "Titolo", "Link"]
    for c, h in enumerate(heads2, 1):
        cell = ws2.cell(row=3, column=c, value=h)
        cell.font, cell.fill = hd, hd_fill
    from common import load_editorial_groups, gruppo_editoriale_di
    groups_map = load_editorial_groups()
    for row, a in enumerate(sorted(esclusi, key=lambda x: x["rank"]), start=4):
        ws2.cell(row=row, column=1, value=a["rank"])
        ws2.cell(row=row, column=2, value=a.get("content_type") or "non_classificato")
        ws2.cell(row=row, column=3, value=a["fonte"])
        ws2.cell(row=row, column=4, value=gruppo_editoriale_di(a["fonte"], groups_map))
        ws2.cell(row=row, column=5, value=a["motivo_esclusione"])
        ws2.cell(row=row, column=6, value=a["title"])
        c7 = ws2.cell(row=row, column=7, value="apri articolo")
        c7.hyperlink = a["link"]
        c7.font = Font(color="0563C1", underline="single")

    out_path = DATA / "composizione.xlsx"
    wb.save(out_path)
    return out_path


def main():
    config = json.load(open(CONFIG_PATH, encoding="utf-8"))
    log(f"STEP 9 — compositore {OSSERVATORIO}: n_articoli_totali={config['n_articoli_totali']}, "
        f"max_per_testata={config['max_per_testata']}, max_per_gruppo_editoriale={config['max_per_gruppo_editoriale']}")

    candidates = load_candidates()
    log(f"  candidati visti dal gate (Si+No, con content_type): {len(candidates)}")

    already_selected = load_already_selected()
    log(f"  già pubblicati in store per {OSSERVATORIO}: {len(already_selected)}")

    selected, esclusi, log_allentamento = select_for_newsletter(candidates, config, already_selected)
    for line in log_allentamento:
        log(f"  {line}")

    out_path = write_output(selected, esclusi, log_allentamento, config)
    log(f"\n  salvato: {out_path}  ({len(selected)} selezionati, {len(esclusi)} esclusi)")
    log("\nFATTO.")


if __name__ == "__main__":
    main()
