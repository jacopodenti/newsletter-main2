"""
STEP 2 — Costruisce il dataset etichettato NUOVO, con TESTO INTERO.

Composizione:
  positivi (1): 26 "Si" del set-50 + 106 editoriali OCX newsletter
  negativi (0): 24 "No" del set-50 + 3 editoriali "non va bene" + 25 palesemente off-topic dal corpus

Ogni riga traccia l'ORIGINE (così puoi filtrare/togliere una fonte quando vuoi).
Scarica il testo intero dagli URL (fallback al sommario se il sito blocca).

Output: data/dataset.jsonl  +  data/dataset_meta.json
Uso:    python ocx_relevance/02_build_dataset.py
"""
import json
import sys
from concurrent.futures import ThreadPoolExecutor

import openpyxl
sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from common import (log, fetch_fulltext, DATA, XLSX_LABEL, XLSX_EDITORIAL)

# parole che segnalano articoli PALESEMENTE fuori tema OCX (per i 25 negativi extra)
OFFTOPIC = ["iphone", "tan ", "taeg", "luce e gas", "prezzo bloccato", "conto a zero",
            "carta gratis", "packaging", "mutuo", "quantum", "crittografia", "truffa informatica",
            "fascicolo sanitario", "cloud sovrano", "school tech", "blue economy", "cleantech"]


def load_50():
    wb = openpyxl.load_workbook(XLSX_LABEL, data_only=True)
    ws = wb["Ranking"]
    # header in riga 2: C=Titolo, D=Sommario, F=Link, H=Rilevante
    rows = []
    for r in range(3, ws.max_row + 1):
        h = ws.cell(row=r, column=8).value
        if not h or str(h).strip() == "":
            continue
        rows.append({
            "title": str(ws.cell(row=r, column=3).value or "").strip(),
            "summary": str(ws.cell(row=r, column=4).value or ""),
            "link": str(ws.cell(row=r, column=6).value or "").strip(),
            "label": 1 if str(h).strip().lower() == "si" else 0,
            "origin": "set50",
        })
    return rows


def load_editorial():
    wb = openpyxl.load_workbook(XLSX_EDITORIAL, data_only=True)
    out, seen = [], set()
    for ws in wb.worksheets:
        hdr = {str(ws.cell(row=1, column=c).value).strip().lower(): c
               for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value}
        tc, sc, lc = hdr.get("titolo"), hdr.get("riassunto"), hdr.get("link")
        nc = next((v for k, v in hdr.items() if k.startswith("note")), None)
        for r in range(2, ws.max_row + 1):
            t = ws.cell(row=r, column=tc).value if tc else None
            if not t or not str(t).strip():
                continue
            link = str(ws.cell(row=r, column=lc).value or "").strip() if lc else ""
            if link and link in seen:
                continue  # stessa notizia già presa da un altro foglio: non contarla due volte
            seen.add(link)
            note = str(ws.cell(row=r, column=nc).value or "").strip().lower() if nc else ""
            # nota con queste parole = l'editor l'ha scartata in fase di redazione -> label negativa
            neg = any(w in note for w in ["non va", "fuori scope", "non bene", "scartato"])
            out.append({
                "title": str(t).strip(),
                "summary": str(ws.cell(row=r, column=sc).value or "") if sc else "",
                "link": link,
                "label": 0 if neg else 1,
                "origin": "editorial",
            })
    return out


def load_offtopic_negatives(exclude_titles, n=25):
    wb = openpyxl.load_workbook(XLSX_LABEL, data_only=True)
    dt = wb["Dati"]
    picks = []
    for r in range(2, dt.max_row + 1):
        t = dt.cell(row=r, column=1).value
        if not t:
            continue
        title = str(t).strip()
        if title in exclude_titles:
            continue
        low = title.lower()
        if any(w in low for w in OFFTOPIC):
            picks.append({
                "title": title,
                "summary": str(dt.cell(row=r, column=17).value or ""),
                "link": str(dt.cell(row=r, column=2).value or "").strip(),
                "label": 0,
                "origin": "offtopic_manual",
            })
        if len(picks) >= n:
            break
    return picks


def main():
    log("STEP 2 — costruzione dataset")
    # 1) carica le tre componenti del dataset e le unisce (l'esclusione evita doppioni
    # tra set-50/editoriali e i negativi off-topic pescati dal corpus)
    d50 = load_50()
    ed = load_editorial()
    exclude = {x["title"] for x in d50} | {x["title"] for x in ed}
    neg_extra = load_offtopic_negatives(exclude, n=25)
    rows = d50 + ed + neg_extra

    log(f"  set-50:            {len(d50):>4}  (Si={sum(x['label'] for x in d50)}, No={sum(1-x['label'] for x in d50)})")
    log(f"  editoriali:        {len(ed):>4}  (Si={sum(x['label'] for x in ed)}, No={sum(1-x['label'] for x in ed)})")
    log(f"  off-topic manuali: {len(neg_extra):>4}  (tutti No)")
    log(f"  TOTALE:            {len(rows):>4}  (positivi={sum(x['label'] for x in rows)}, negativi={sum(1-x['label'] for x in rows)})")

    # 2) scarica il testo intero dagli URL (fallback al sommario se il sito blocca), in parallelo
    log("\n  scarico il testo intero dagli URL...")
    def fetch(x):
        text, src = fetch_fulltext(x["link"], x["summary"])
        x["full_text"] = text
        x["text_source"] = src
        return x
    with ThreadPoolExecutor(max_workers=16) as ex:
        rows = list(ex.map(fetch, rows))
    nfull = sum(1 for x in rows if x["text_source"] == "full")
    log(f"  testo intero: {nfull}/{len(rows)} ({nfull/len(rows):.0%})  fallback sommario: {len(rows)-nfull}")

    # 3) salva dataset.jsonl + dataset_meta.json
    with open(DATA / "dataset.jsonl", "w", encoding="utf-8") as f:
        for i, x in enumerate(rows):
            x["id"] = i
            f.write(json.dumps(x, ensure_ascii=False) + "\n")
    meta = {
        "totale": len(rows),
        "positivi": sum(x["label"] for x in rows),
        "negativi": sum(1 - x["label"] for x in rows),
        "per_origine": {o: sum(1 for x in rows if x["origin"] == o) for o in {x["origin"] for x in rows}},
        "testo_intero": nfull, "fallback_sommario": len(rows) - nfull,
    }
    json.dump(meta, open(DATA / "dataset_meta.json", "w"), ensure_ascii=False, indent=2)
    log(f"\n  salvato: {DATA/'dataset.jsonl'}  ({len(rows)} righe)")
    log(f"  salvato: {DATA/'dataset_meta.json'}")
    # 4) stampa i titoli off-topic scelti, per un controllo manuale rapido
    log("\n  ispeziona i 25 negativi off-topic scelti (controlla che siano davvero fuori tema):")
    for x in neg_extra:
        log(f"    - {x['title'][:70]}")
    log("\nFATTO.")


if __name__ == "__main__":
    main()
