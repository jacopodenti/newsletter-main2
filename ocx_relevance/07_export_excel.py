"""
STEP 7 — Scrive il ranking come nuovo foglio DENTRO LABEL-Newsletter.xlsx.

Separato da 06_rank.py così si può rigenerare/aggiornare il foglio senza
rifare embedding o chiamate gate. STESSO SCHEMA di ai_relevance/07_export_excel.py
e startup_thinking_relevance/07_export_excel.py.

Non rifà nessuna chiamata API: legge solo i file già prodotti (ranking.csv,
corpus_fulltext_cache.json).

Output: LABEL-Newsletter.xlsx -> nuovo foglio "Ranking_OCX_nuovo"
Uso:    python ocx_relevance/07_export_excel.py
"""
import argparse
import csv
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DATA = HERE / "data"
MODELS = HERE / "models"
XLSX_LABEL = ROOT / "LABEL-Newsletter.xlsx"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ranking", default="ranking.csv", help="file ranking da usare (dentro data/)")
    ap.add_argument("--sheet", default="Ranking_OCX_nuovo", help="nome del foglio da scrivere")
    args = ap.parse_args()

    # 1) carica ranking.csv + cache testo (per sapere se un articolo ha testo completo o solo sommario)
    ranking = list(csv.DictReader(open(DATA / args.ranking, encoding="utf-8")))
    cache = json.load(open(MODELS / "corpus_fulltext_cache.json", encoding="utf-8"))

    si = sum(1 for r in ranking if r["pertinenza_llm"] == "Si")
    topn = sum(1 for r in ranking if r["pertinenza_llm"] in ("Si", "No"))
    n = len(ranking)

    # 2) crea il foglio (lo ricrea da zero se esiste già, per poterlo rigenerare)
    SHEET = args.sheet
    wb = openpyxl.load_workbook(XLSX_LABEL)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)

    hf = Font(bold=True, color="FFFFFF")
    hd = PatternFill("solid", fgColor="1F6F3E")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    orange = PatternFill("solid", fgColor="FFE0B3")
    # Titolo colorato = sospetto duplicato (embedding molto simile a un altro candidato
    # dello stesso run, es. stessa notizia da fonti diverse). ATTENZIONE (audit 22/08): la
    # palette ha solo 8 colori, i cluster reali possono superarli abbondantemente (153 nel run
    # del 21/08) — dal nono cluster in poi il colore si ripete su gruppi scollegati. Il colore
    # resta solo un aiuto visivo per un'occhiata rapida; il numero di cluster scritto accanto al
    # titolo (sempre univoco) è l'informazione affidabile, non il colore.
    CLUSTER_COLORS = ["FFD9D9", "D9E8FF", "D9FFE0", "F0D9FF", "D9FFF7", "FFE0F0", "E8FFD9", "E0E0FF"]

    def cluster_fill(cluster_id):
        if not cluster_id:
            return None
        return PatternFill("solid", fgColor=CLUSTER_COLORS[int(cluster_id) % len(CLUSTER_COLORS)])

    ws["A1"] = (f"Ranking OCX (ranker + gate LLM su top-{topn}). Rosa finale: {si}.  "
                f"Parere ricercatore: da compilare a mano (Si/No).  Testo: giallo = solo sommario RSS, "
                f"arancione = estratto parziale (il resto dell'articolo era testo offuscato/non affidabile, "
                f"vedi audit 24/08 — valutato su un frammento, non sull'articolo intero).  "
                f"Titolo con [cluster N]: sospetto duplicato (stessa notizia da fonti diverse) — il numero identifica il "
                f"gruppo in modo univoco, il colore è solo un aiuto visivo e si ripete oltre 8 cluster nello stesso run "
                f"(fidarsi del numero, non del colore, se sono più di 8). Vedi motivo del revisore LLM.")
    ws["A1"].font = Font(bold=True, size=12)

    # 3) intestazioni
    heads = ["Rank", "Score ranker", "Pertinenza LLM", "Parere ricercatore", "Testo", "Titolo", "Fonte", "Data", "Link", "Parere revisore LLM", "Confidence LLM"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = hf
        cell.fill = hd

    # 4) scrive le righe: una per articolo, con evidenziazione testo-mancante e cluster duplicati
    for i, r in enumerate(ranking):
        row = 3 + i + 1
        title = r["titolo"]
        text_source = cache.get(title, {}).get("source")

        ws.cell(row=row, column=1, value=int(r["rank"]))
        ws.cell(row=row, column=2, value=float(r["score_ranker"]))
        ws.cell(row=row, column=3, value=r["pertinenza_llm"])
        # colonna 4 "Parere ricercatore" lasciata vuota apposta: da compilare a mano (dropdown Si/No)
        # text_source: 'full' = completo, 'partial' = estratto parziale (testo offuscato tagliato
        # al punto di degrado, vedi audit 24/08), tutto il resto = solo il summary RSS.
        if text_source == "full":
            c5 = ws.cell(row=row, column=5, value="completo")
        elif text_source == "partial":
            c5 = ws.cell(row=row, column=5, value="estratto parziale")
            c5.fill = orange
        else:
            c5 = ws.cell(row=row, column=5, value="solo sommario")
            c5.fill = yellow
        cid = r.get("cluster_id")
        c6 = ws.cell(row=row, column=6, value=f"{title}  [cluster {cid}]" if cid else title)
        cf = cluster_fill(cid)
        if cf:
            c6.fill = cf
        ws.cell(row=row, column=7, value=r["fonte"])
        ws.cell(row=row, column=8, value=r.get("data", ""))
        c9 = ws.cell(row=row, column=9, value="apri articolo" if r["link"] else "")
        if r["link"]:
            c9.hyperlink = r["link"]
            c9.font = Font(color="0563C1", underline="single")
        ws.cell(row=row, column=10, value=r.get("motivo_llm", ""))
        conf = r.get("confidence_llm", "")
        ws.cell(row=row, column=11, value=float(conf) if conf not in ("", None) else "")

    # 5) dropdown Si/No per il parere del ricercatore (colonna D)
    last_row = 3 + n
    dv = DataValidation(type="list", formula1='"Si,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D4:D{last_row}")

    # 6) formattazione finale (larghezza colonne, riquadro fisso, filtro) e salvataggio
    widths = {"A": 6, "B": 12, "C": 14, "D": 16, "E": 13, "F": 60, "G": 20, "H": 22, "I": 14, "J": 55, "K": 12}
    for col, wd in widths.items():
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:K{2 + n + 1}"

    wb.save(XLSX_LABEL)
    print(f"salvato foglio '{SHEET}' in {XLSX_LABEL.name}  ({n} righe)")


if __name__ == "__main__":
    main()
