"""
Test decisivo: GPT-5.4 sui 110 editoriali (tutti positivi umani).
+ overlap con il corpus 1418. + confronto prompt severo vs morbido.
"""
import os, sys, json, threading
from concurrent.futures import ThreadPoolExecutor, as_completed
sys.path.insert(0, "src")
from dotenv import load_dotenv; load_dotenv("/Users/gabriele/Desktop/newsletter-main2/.env")
import openpyxl
from openai import OpenAI
ROOT = "/Users/gabriele/Desktop/newsletter-main2"
client = OpenAI(api_key=os.environ["AZURE_OPENAI_API_KEY"], base_url=os.environ["AZURE_OPENAI_ENDPOINT"])

oss = json.load(open(f"{ROOT}/config/osservatori.json"))
ocx = next(o for o in oss if o["slug"] == "omnichannel-customer-experience")
desc = ocx.get("descrizione_enhanced") or ocx["descrizione"]

STRICT = f"""Sei un classificatore di pertinenza per l'Osservatorio Omnichannel Customer Experience (OCX).
DESCRIZIONE:\n{desc}
E pertinente SOLO se tratta customer experience omnicanale, customer journey, contact center, CRM,
self-service, loyalty, NPS. NON e pertinente se parla genericamente di digitale, AI, cybersecurity,
cloud, innovazione, startup, fintech, pagamenti, telco, hardware senza focus sull'esperienza cliente.
Rispondi SOLO JSON: {{"rilevante": true/false, "motivo": "frase"}}"""

SOFT = f"""Sei un analista dell'Osservatorio Omnichannel Customer Experience (OCX).
DESCRIZIONE:\n{desc}
Decidi se l'articolo e RILEVANTE per un lettore di questo osservatorio: temi di customer experience,
relazione col cliente, canali, marketing/vendite/servizio orientati al cliente, tecnologie e strategie
che impattano l'esperienza del cliente. Sii inclusivo: se tocca l'esperienza cliente anche in modo
applicato (retail, banche, AI nel servizio, loyalty, dati cliente), e rilevante.
Rispondi SOLO JSON: {{"rilevante": true/false, "motivo": "frase"}}"""

# --- 110 editoriali ---
wb = openpyxl.load_workbook(f"{ROOT}/OCX26_Newsletter.xlsx", data_only=True)
picks, seen = [], set()
for ws in wb.worksheets:
    hdr = {str(ws.cell(row=1, column=c).value).strip().lower(): c for c in range(1, ws.max_column+1) if ws.cell(row=1, column=c).value}
    tc, sc, lc = hdr.get("titolo"), hdr.get("riassunto"), hdr.get("link")
    for r in range(2, ws.max_row+1):
        t = ws.cell(row=r, column=tc).value if tc else None
        if not t or not str(t).strip(): continue
        link = str(ws.cell(row=r, column=lc).value or "") if lc else ""
        if link and link in seen: continue
        seen.add(link)
        picks.append({"title": str(t), "summary": str(ws.cell(row=r, column=sc).value or "") if sc else "", "link": link})

# --- overlap col corpus 1418 ---
dt = openpyxl.load_workbook(f"{ROOT}/LABEL-Newsletter.xlsx", data_only=True)["Dati"]
corpus_links = {str(dt.cell(row=r, column=2).value or "").strip() for r in range(2, 1420)}
corpus_titles = {str(dt.cell(row=r, column=1).value or "").strip().lower() for r in range(2, 1420)}
ov = sum(1 for p in picks if p["link"].strip() in corpus_links or p["title"].strip().lower() in corpus_titles)
print(f"Editoriali: {len(picks)} | nel corpus 1418: {ov}")

def run(instr, tag):
    def one(p):
        prompt = f"{instr}\n\nARTICOLO:\nTitolo: {p['title']}\nSommario: {p['summary']}\n\nJSON:"
        for _ in range(3):
            try:
                r = client.responses.create(model="gpt-5.4", input=prompt)
                txt = r.output_text.strip()
                if txt.startswith("```"): txt = txt.split("```")[1].removeprefix("json").strip()
                return 1 if json.loads(txt).get("rilevante") else 0
            except Exception:
                continue
        return None
    with ThreadPoolExecutor(max_workers=12) as ex:
        res = list(ex.map(one, picks))
    si = sum(1 for x in res if x == 1); no = sum(1 for x in res if x == 0); err = sum(1 for x in res if x is None)
    print(f"[{tag}] GPT sui 110 editoriali: Si={si} ({si/len(picks):.0%})  No={no}  err={err}")
    return res

print("\n--- PROMPT SEVERO (quello usato sui 1418) ---")
rs = run(STRICT, "severo")
print("--- PROMPT MORBIDO (inclusivo) ---")
rm = run(SOFT, "morbido")

print("\nEsempi bocciati dal SEVERO (No su articoli scelti dai redattori):")
n = 0
for p, x in zip(picks, rs):
    if x == 0 and n < 12:
        print(f"  - {p['title'][:68]}"); n += 1
