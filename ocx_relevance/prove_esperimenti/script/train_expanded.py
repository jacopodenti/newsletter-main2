"""
(1) Training set espanso: 106 positivi editoriali + 26 Si + 27 negativi.
Riaddestra kNN, ricalibra soglia, k-fold onesto.
"""
import sys, pickle
sys.path.insert(0, "src")
from dotenv import load_dotenv; load_dotenv("/Users/gabriele/Desktop/newsletter-main2/.env")
import numpy as np
import openpyxl
from sklearn.model_selection import RepeatedStratifiedKFold
from sklearn.metrics import roc_auc_score
from sklearn.linear_model import LogisticRegression
import embed_classify as ec
from knn_relevance import KNNRelevanceScorer, load_labeled_set, clean_threshold

client = ec._get_client()

def embed(texts):
    out = []
    for i in range(0, len(texts), 16):
        out.extend(ec._embed(client, texts[i:i+16]))
    return np.array(out)

# --- positivi editoriali (newsletter) ---
wb = openpyxl.load_workbook("OCX26_Newsletter.xlsx", data_only=True)
edit_pos, edit_neg, seen = [], [], set()
for ws in wb.worksheets:
    hdr = {str(ws.cell(row=1, column=c).value).strip().lower(): c
           for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value}
    tc, sc, lc = hdr.get("titolo"), hdr.get("riassunto"), hdr.get("link")
    nc = next((v for k, v in hdr.items() if k.startswith("note")), None)
    for r in range(2, ws.max_row + 1):
        t = ws.cell(row=r, column=tc).value if tc else None
        if not t or not str(t).strip(): continue
        link = str(ws.cell(row=r, column=lc).value or "") if lc else ""
        if link and link in seen: continue
        seen.add(link)
        note = str(ws.cell(row=r, column=nc).value or "").strip().lower() if nc else ""
        txt = ec.normalize_text(f"{t}. {ws.cell(row=r, column=sc).value or '' if sc else ''}")
        (edit_neg if any(w in note for w in ["non va","fuori scope","non bene","scartato"]) else edit_pos).append(txt)

# --- 50 etichettati ---
rows50 = load_labeled_set("LABEL-Newsletter.backup.xlsx", sheet="Ranking")
t50 = [ec.normalize_text(f"{r['title']}. {r['summary']}") for r in rows50]
y50 = [r["label"] for r in rows50]

# --- dataset combinato ---
texts = edit_pos + edit_neg + t50
y = np.array([1]*len(edit_pos) + [0]*len(edit_neg) + y50)
print(f"Training set: {len(y)}  positivi={int(y.sum())}  negativi={int((1-y).sum())}")
print(f"  (editoriali +{len(edit_pos)}pos/{len(edit_neg)}neg, 50-set +{sum(y50)}pos/{len(y50)-sum(y50)}neg)")

X = embed(texts)

# --- k-fold onesto (kNN vs LogReg) ---
rkf = RepeatedStratifiedKFold(n_splits=5, n_repeats=20, random_state=0)
splits = list(rkf.split(X, y))
def knn_oof():
    aucs=[]
    for rep in range(20):
        oof=np.zeros(len(y))
        for f in range(5):
            tr,te=splits[rep*5+f]
            oof[te]=KNNRelevanceScorer(k=3).fit(X[tr],y[tr]).score_batch(X[te])
        aucs.append(roc_auc_score(y,oof))
    return np.array(aucs)
def lr_oof():
    aucs=[]
    for rep in range(20):
        oof=np.zeros(len(y))
        for f in range(5):
            tr,te=splits[rep*5+f]
            oof[te]=LogisticRegression(C=1,max_iter=2000).fit(X[tr],y[tr]).predict_proba(X[te])[:,1]
        aucs.append(roc_auc_score(y,oof))
    return np.array(aucs)
ka, la = knn_oof(), lr_oof()
print(f"\nk-fold (5x20) AUC:")
print(f"  kNN k=3   : {ka.mean():.3f} +/- {ka.std():.3f}")
print(f"  LogReg    : {la.mean():.3f} +/- {la.std():.3f}")

# --- soglia pulita ricalibrata (LOO sul combinato) ---
loo=np.zeros(len(y))
for i in range(len(y)):
    m=np.ones(len(y),bool); m[i]=False
    loo[i]=KNNRelevanceScorer(k=3).fit(X[m],y[m]).score(X[i])
thr=clean_threshold(loo,y,margin_frac=0.5)
above=(loo>=thr)
print(f"\nSoglia pulita RICALIBRATA: {thr:.3f}  (prima era 0.143 sui soli 50)")
print(f"  sopra soglia: {above.sum()} articoli, di cui positivi {int(y[above].sum())}/{int(above.sum())}")

# salva vettori+label per riuso
np.savez_compressed("/Users/gabriele/Desktop/newsletter-main2/models/train_expanded.npz", X=X, y=y)
print("\nSalvato models/train_expanded.npz (vettori + label del set espanso)")
