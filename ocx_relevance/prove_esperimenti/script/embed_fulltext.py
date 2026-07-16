"""
Confronto embedding: SOMMARIO vs TESTO INTERO.
- AUC leave-one-out sui 50 etichettati (head-to-head).
- Ri-embedda tutto il corpus sul testo intero -> corpus_vectors_full.npz.
"""
import sys, json
sys.path.insert(0, "src")
from dotenv import load_dotenv; load_dotenv("/Users/gabriele/Desktop/newsletter-main2/.env")
import numpy as np
import openpyxl
from sklearn.metrics import roc_auc_score
import embed_classify as ec
from knn_relevance import KNNRelevanceScorer, load_labeled_set
ROOT = "/Users/gabriele/Desktop/newsletter-main2"
ft = json.load(open(f"{ROOT}/models/fulltext.json"))

def embed(texts):
    out = []
    for i in range(0, len(texts), 16):
        out.extend(ec._embed(ec._get_client(), texts[i:i+16]))
    return np.array(out)

# --- 50 etichettati: sommario vs testo intero ---
rows = load_labeled_set("LABEL-Newsletter.backup.xlsx", sheet="Ranking")
y = np.array([r["label"] for r in rows])
X_sum = embed([ec.normalize_text(f"{r['title']}. {r['summary']}") for r in rows])
X_full = embed([ec.normalize_text(ft.get(r["title"].strip(), {}).get("text", f"{r['title']}. {r['summary']}")) for r in rows])
nfull = sum(1 for r in rows if ft.get(r["title"].strip(), {}).get("source") == "full")
print(f"50 etichettati: {nfull} con testo intero, {50-nfull} fallback sommario")

def loo_auc(X):
    s = np.zeros(len(y))
    for i in range(len(y)):
        m = np.ones(len(y), bool); m[i] = False
        s[i] = KNNRelevanceScorer(k=3).fit(X[m], y[m]).score(X[i])
    return roc_auc_score(y, s)

print(f"\n=== AUC kNN leave-one-out (50 etichettati) ===")
print(f"  SOMMARIO (titolo+25 parole): {loo_auc(X_sum):.3f}")
print(f"  TESTO INTERO:                {loo_auc(X_full):.3f}")

# --- ri-embedda tutto il corpus sul testo intero ---
dt = openpyxl.load_workbook(f"{ROOT}/LABEL-Newsletter.xlsx", data_only=True)["Dati"]
titles = [str(dt.cell(row=d, column=1).value).strip() for d in range(2, 1420) if dt.cell(row=d, column=1).value]
texts = [ec.normalize_text(ft.get(t, {}).get("text", t)) for t in titles]
print(f"\nRi-embeddo {len(texts)} articoli sul testo intero...")
Vf = []
for i in range(0, len(texts), 16):
    Vf.extend(ec._embed(ec._get_client(), texts[i:i+16]))
    if (i // 16) % 15 == 0: print(f"  {i}/{len(texts)}", flush=True)
Vf = np.array(Vf)
np.savez_compressed(f"{ROOT}/models/corpus_vectors_full.npz", V=Vf)
print(f"Salvato models/corpus_vectors_full.npz")
print("FINITO", flush=True)
