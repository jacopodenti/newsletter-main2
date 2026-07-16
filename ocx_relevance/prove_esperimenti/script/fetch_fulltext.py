"""
Scarica il testo INTERO dei 1418 articoli dagli URL (colonna LINK).
Estrae il corpo con trafilatura, fallback al sommario se fallisce.
Salva models/fulltext.json.
"""
import os, sys, json, threading
from concurrent.futures import ThreadPoolExecutor
import requests
import trafilatura
import openpyxl
ROOT = "/Users/gabriele/Desktop/newsletter-main2"

dt = openpyxl.load_workbook(f"{ROOT}/LABEL-Newsletter.xlsx", data_only=True)["Dati"]
arts = []
for d in range(2, 1420):
    t = dt.cell(row=d, column=1).value
    if not t: continue
    arts.append({"title": str(t).strip(),
                 "link": str(dt.cell(row=d, column=2).value or "").strip(),
                 "summary": str(dt.cell(row=d, column=17).value or "")})
N = len(arts)
print(f"Articoli: {N}", flush=True)

UA = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"}
lock = threading.Lock(); ok = fail = 0

def fetch(a):
    global ok, fail
    link = a["link"]
    if not link.startswith("http"):
        with lock: fail += 1
        return a["title"], {"text": a["summary"], "source": "summary", "chars": len(a["summary"])}
    try:
        r = requests.get(link, headers=UA, timeout=15)
        body = trafilatura.extract(r.text, include_comments=False, include_tables=False, favor_precision=True) if r.ok else None
    except Exception:
        body = None
    if body and len(body) > len(a["summary"]) + 50:
        with lock: ok += 1
        return a["title"], {"text": body[:8000], "source": "full", "chars": len(body)}
    else:
        with lock: fail += 1
        return a["title"], {"text": a["summary"], "source": "summary", "chars": len(a["summary"])}

out = {}
done = 0
with ThreadPoolExecutor(max_workers=16) as ex:
    for title, rec in ex.map(fetch, arts):
        out[title] = rec; done += 1
        if done % 100 == 0: print(f"  {done}/{N}  (full={ok} fallback={fail})", flush=True)

json.dump(out, open(f"{ROOT}/models/fulltext.json", "w"), ensure_ascii=False)
full = [v for v in out.values() if v["source"] == "full"]
fb = [v for v in out.values() if v["source"] == "summary"]
import statistics
print(f"\n=== RISULTATO ===")
print(f"  testo intero estratto: {len(full)}/{N} ({len(full)/N:.0%})")
print(f"  fallback al sommario:  {len(fb)}/{N}")
if full:
    print(f"  lunghezza testo intero: mediana {int(statistics.median([v['chars'] for v in full]))} car "
          f"(vs sommario ~153) -> ~{int(statistics.median([v['chars'] for v in full])/6)} parole")
print("FINITO", flush=True)
