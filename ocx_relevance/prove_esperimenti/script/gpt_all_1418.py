"""
GPT-5.4 previsione pertinenza Si/No su TUTTI i 1418 (parallelo).
Aggiunge la colonna 'Pertinenza LLM' al foglio Ranking_kNN_v2.
"""
import os, sys, json, threading
from concurrent.futures import ThreadPoolExecutor, as_completed
sys.path.insert(0, "src")
from dotenv import load_dotenv; load_dotenv("/Users/gabriele/Desktop/newsletter-main2/.env")
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openai import OpenAI
import embed_classify as ec
from knn_relevance import KNNRelevanceScorer, load_labeled_set, clean_threshold

ROOT = "/Users/gabriele/Desktop/newsletter-main2"
MODEL = "gpt-5.4"
client = OpenAI(api_key=os.environ["AZURE_OPENAI_API_KEY"], base_url=os.environ["AZURE_OPENAI_ENDPOINT"])

oss = json.load(open(f"{ROOT}/config/osservatori.json"))
ocx_o = next(o for o in oss if o["slug"] == "omnichannel-customer-experience")
ocx_desc = ocx_o.get("descrizione_enhanced") or ocx_o["descrizione"]
INSTR = f"""Sei un classificatore di pertinenza per l'Osservatorio "Omnichannel Customer Experience" (OCX) del Politecnico di Milano.

DESCRIZIONE OSSERVATORIO:
{ocx_desc}

COMPITO: dato un articolo (titolo + sommario), decidi se e PERTINENTE a questo osservatorio.
E pertinente SOLO se tratta in modo specifico di customer experience omnicanale, customer journey,
gestione della relazione col cliente su piu canali, contact center, CRM, self-service, loyalty, NPS.
NON e pertinente se parla genericamente di digitale, AI, cybersecurity, cloud, innovazione, startup,
fintech, pagamenti, telco, hardware senza un focus concreto sull'esperienza del cliente.

Rispondi SOLO con JSON: {{"rilevante": true/false, "confidence": 0.0-1.0, "motivo": "una frase"}}"""

# --- corpus articoli (ordine Dati) ---
dt = openpyxl.load_workbook(f"{ROOT}/LABEL-Newsletter.xlsx", data_only=True)["Dati"]
arts = []
for d in range(2, 1420):
    t = dt.cell(row=d, column=1).value
    if not t: continue
    arts.append({"title": str(t).strip(), "summary": str(dt.cell(row=d, column=17).value or ""),
                 "fonte": str(dt.cell(row=d, column=4).value or ""), "data": str(dt.cell(row=d, column=3).value or ""),
                 "link": str(dt.cell(row=d, column=2).value or ""), "ct": str(dt.cell(row=d, column=6).value or "")})
N = len(arts)
print(f"Articoli: {N} | modello {MODEL}")

lock = threading.Lock()
tok_in = tok_out = 0
def classify(i):
    global tok_in, tok_out
    a = arts[i]
    prompt = f"{INSTR}\n\nARTICOLO:\nTitolo: {a['title']}\nSommario: {a['summary']}\n\nJSON:"
    for attempt in range(3):
        try:
            r = client.responses.create(model=MODEL, input=prompt)
            txt = r.output_text.strip()
            if txt.startswith("```"): txt = txt.split("```")[1].removeprefix("json").strip()
            res = json.loads(txt)
            with lock:
                tok_in += r.usage.input_tokens; tok_out += r.usage.output_tokens
            return i, (1 if res.get("rilevante") else 0)
        except Exception as e:
            if attempt == 2:
                return i, None
    return i, None

preds = [None] * N
done = 0
with ThreadPoolExecutor(max_workers=12) as ex:
    futs = [ex.submit(classify, i) for i in range(N)]
    for f in as_completed(futs):
        i, p = f.result(); preds[i] = p
        done += 1
        if done % 100 == 0: print(f"  {done}/{N}...", flush=True)

si = sum(1 for p in preds if p == 1); no = sum(1 for p in preds if p == 0); err = sum(1 for p in preds if p is None)
print(f"\nLLM: Si={si}  No={no}  errori={err}")
json.dump({arts[i]["title"]: preds[i] for i in range(N)}, open(f"{ROOT}/models/gpt_relevance.json", "w"), ensure_ascii=False)
print(f"token: {tok_in} in + {tok_out} out")

# --- kNN score + ordine ---
V = np.load(f"{ROOT}/models/corpus_vectors.npz")["V"]
tr = np.load(f"{ROOT}/models/train_expanded.npz"); knn = KNNRelevanceScorer(k=3).fit(tr["X"], tr["y"])
sc = knn.score_batch(V); order = np.argsort(-sc)
loo = np.zeros(len(tr["y"]))
for i in range(len(tr["y"])):
    m = np.ones(len(tr["y"]), bool); m[i] = False
    loo[i] = KNNRelevanceScorer(k=3).fit(tr["X"][m], tr["y"][m]).score(tr["X"][i])
thr = clean_threshold(loo, tr["y"], 0.5)
lab = {r["title"].strip(): r["label"] for r in load_labeled_set("LABEL-Newsletter.backup.xlsx", sheet="Ranking")}

# --- incrocio kNN vs LLM ---
above = sc >= thr
llm_si = np.array([preds[i] == 1 for i in range(N)])
print("\n=== INCROCIO kNN-soglia vs LLM ===")
print(f"  sopra soglia kNN: {above.sum()} | di cui LLM-Si: {(above & llm_si).sum()}")
print(f"  LLM-Si totali: {llm_si.sum()} | di cui sotto soglia kNN: {(~above & llm_si).sum()}")

# --- riscrivi foglio con colonna LLM ---
wb = openpyxl.load_workbook(f"{ROOT}/LABEL-Newsletter.xlsx")
if "Ranking_kNN_v2" in wb.sheetnames: del wb["Ranking_kNN_v2"]
ws = wb.create_sheet("Ranking_kNN_v2")
hf = Font(bold=True, color="FFFFFF"); hd = PatternFill("solid", fgColor="1F6F3E")
green = PatternFill("solid", fgColor="E2EFDA"); red = PatternFill("solid", fgColor="FCE4D6")
yell = PatternFill("solid", fgColor="FFF2CC")
ws["A1"] = "Ranking kNN v2 + Pertinenza LLM (gpt-5.4) su tutti i 1418"
ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = f"Soglia kNN {thr:.3f}. LLM Si={si}/No={no}. Verde=Si noto, Rosso=No noto."
ws["A2"].font = Font(italic=True, size=10, color="666666")
heads = ["Rank","Score kNN","Sopra soglia kNN","Pertinenza LLM","Label nota","Titolo","Sommario","Fonte","Data","Content Type","Link"]
for c, h in enumerate(heads, 1):
    cell = ws.cell(row=4, column=c, value=h); cell.font = hf; cell.fill = hd
for rank, i in enumerate(order, 1):
    a = arts[i]; row = 4 + rank; L = lab.get(a["title"]); pl = preds[i]
    ws.cell(row=row, column=1, value=rank)
    ws.cell(row=row, column=2, value=round(float(sc[i]), 4))
    ws.cell(row=row, column=3, value="SI" if sc[i] >= thr else "")
    c4 = ws.cell(row=row, column=4, value=("Si" if pl == 1 else "No") if pl is not None else "?")
    if pl == 1: c4.fill = yell
    ws.cell(row=row, column=5, value=("Si" if L == 1 else "No") if L is not None else "")
    ws.cell(row=row, column=6, value=a["title"])
    ws.cell(row=row, column=7, value=a["summary"][:400])
    ws.cell(row=row, column=8, value=a["fonte"]); ws.cell(row=row, column=9, value=a["data"])
    ws.cell(row=row, column=10, value=a["ct"]); ws.cell(row=row, column=11, value=a["link"])
    if L == 0:
        for c in range(1, 12): ws.cell(row=row, column=c).fill = red
    elif L == 1:
        for c in range(1, 12): ws.cell(row=row, column=c).fill = green
    ws.cell(row=row, column=7).alignment = Alignment(wrap_text=True, vertical="top")
for col, w in {"A":6,"B":10,"C":15,"D":14,"E":10,"F":55,"G":58,"H":18,"I":20,"J":18,"K":26}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
wb.save(f"{ROOT}/LABEL-Newsletter.xlsx")
print("Foglio 'Ranking_kNN_v2' aggiornato con colonna Pertinenza LLM.")
