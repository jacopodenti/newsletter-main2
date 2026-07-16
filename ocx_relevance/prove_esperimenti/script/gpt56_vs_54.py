"""
gpt-5.6-sol su 1418 (stesso prompt severo del run 5.4) + confronto con 5.4 cached.
"""
import os, sys, json, threading
from concurrent.futures import ThreadPoolExecutor
sys.path.insert(0, "src")
from dotenv import load_dotenv; load_dotenv("/Users/gabriele/Desktop/newsletter-main2/.env")
import openpyxl
from openai import OpenAI
from knn_relevance import load_labeled_set
ROOT = "/Users/gabriele/Desktop/newsletter-main2"
MODEL = "gpt-5.6-sol"
client = OpenAI(api_key=os.environ["AZURE_OPENAI_API_KEY"], base_url=os.environ["AZURE_OPENAI_ENDPOINT"])

oss = json.load(open(f"{ROOT}/config/osservatori.json"))
ocx = next(o for o in oss if o["slug"] == "omnichannel-customer-experience")
desc = ocx.get("descrizione_enhanced") or ocx["descrizione"]
INSTR = f"""Sei un classificatore di pertinenza per l'Osservatorio "Omnichannel Customer Experience" (OCX) del Politecnico di Milano.

DESCRIZIONE OSSERVATORIO:
{desc}

COMPITO: dato un articolo (titolo + sommario), decidi se e PERTINENTE a questo osservatorio.
E pertinente SOLO se tratta in modo specifico di customer experience omnicanale, customer journey,
gestione della relazione col cliente su piu canali, contact center, CRM, self-service, loyalty, NPS.
NON e pertinente se parla genericamente di digitale, AI, cybersecurity, cloud, innovazione, startup,
fintech, pagamenti, telco, hardware senza un focus concreto sull'esperienza del cliente.

Rispondi SOLO con JSON: {{"rilevante": true/false, "confidence": 0.0-1.0, "motivo": "una frase"}}"""

dt = openpyxl.load_workbook(f"{ROOT}/LABEL-Newsletter.xlsx", data_only=True)["Dati"]
arts = []
for d in range(2, 1420):
    t = dt.cell(row=d, column=1).value
    if not t: continue
    arts.append({"title": str(t).strip(), "summary": str(dt.cell(row=d, column=17).value or "")})
N = len(arts)
print(f"Articoli: {N} | modello {MODEL}", flush=True)

lock = threading.Lock(); tin = tout = 0
def classify(i):
    global tin, tout
    a = arts[i]
    prompt = f"{INSTR}\n\nARTICOLO:\nTitolo: {a['title']}\nSommario: {a['summary']}\n\nJSON:"
    for _ in range(3):
        try:
            r = client.responses.create(model=MODEL, input=prompt)
            txt = r.output_text.strip()
            if txt.startswith("```"): txt = txt.split("```")[1].removeprefix("json").strip()
            res = json.loads(txt)
            with lock:
                tin += r.usage.input_tokens; tout += r.usage.output_tokens
            return i, (1 if res.get("rilevante") else 0)
        except Exception:
            continue
    return i, None

preds = [None] * N; done = 0
with ThreadPoolExecutor(max_workers=12) as ex:
    for i, p in ex.map(lambda i: classify(i), range(N)):
        preds[i] = p; done += 1
        if done % 100 == 0: print(f"  {done}/{N}", flush=True)

json.dump({arts[i]["title"]: preds[i] for i in range(N)}, open(f"{ROOT}/models/gpt56_relevance.json", "w"), ensure_ascii=False)
si6 = sum(1 for p in preds if p == 1); no6 = sum(1 for p in preds if p == 0); err6 = sum(1 for p in preds if p is None)
print(f"\n[5.6-sol] Si={si6} No={no6} err={err6} | token {tin} in + {tout} out", flush=True)

# --- confronto con 5.4 cached ---
g54 = json.load(open(f"{ROOT}/models/gpt_relevance.json"))
si4 = sum(1 for t in g54.values() if t == 1)
both = same = onlyt6 = onlyt4 = 0
for a in arts:
    p6 = preds[arts.index(a)] if False else None
# indicizza per titolo
p6map = {arts[i]["title"]: preds[i] for i in range(N)}
agree = dis = b_si = only6 = only4 = 0
for t in p6map:
    x6 = p6map[t]; x4 = g54.get(t)
    if x6 is None or x4 is None: continue
    if x6 == x4: agree += 1
    else: dis += 1
    if x6 == 1 and x4 == 1: b_si += 1
    if x6 == 1 and x4 == 0: only6 += 1
    if x6 == 0 and x4 == 1: only4 += 1
print("\n=== CONFRONTO 5.6-sol vs 5.4 ===")
print(f"  Si: 5.6-sol={si6}  |  5.4={si4}")
print(f"  accordo: {agree}  disaccordo: {dis}")
print(f"  entrambi Si: {b_si} | solo 5.6-sol Si: {only6} | solo 5.4 Si: {only4}")

# precisione su label note
lab = {r["title"].strip(): r["label"] for r in load_labeled_set("LABEL-Newsletter.backup.xlsx", sheet="Ranking")}
for name, pm in [("5.6-sol", p6map), ("5.4", g54)]:
    tp = sum(1 for t, l in lab.items() if pm.get(t) == 1 and l == 1)
    fp = sum(1 for t, l in lab.items() if pm.get(t) == 1 and l == 0)
    fn = sum(1 for t, l in lab.items() if pm.get(t) == 0 and l == 1)
    print(f"  [{name}] su 50 note: TP={tp} FP={fp} FN={fn}  precisione={tp/(tp+fp) if tp+fp else 0:.0%} recall={tp/(tp+fn) if tp+fn else 0:.0%}")

# --- foglio confronto (ordinato per kNN) ---
import numpy as np
from openpyxl.styles import Font, PatternFill, Alignment
from knn_relevance import KNNRelevanceScorer
V = np.load(f"{ROOT}/models/corpus_vectors.npz")["V"]
tr = np.load(f"{ROOT}/models/train_expanded.npz")
sc = KNNRelevanceScorer(k=3).fit(tr["X"], tr["y"]).score_batch(V)
order = np.argsort(-sc)
dtf = openpyxl.load_workbook(f"{ROOT}/LABEL-Newsletter.xlsx", data_only=True)["Dati"]
meta = {}
for d in range(2, 1420):
    t = dtf.cell(row=d, column=1).value
    if not t: continue
    meta[str(t).strip()] = (str(dtf.cell(row=d, column=4).value or ""), str(dtf.cell(row=d, column=2).value or ""))

wb = openpyxl.load_workbook(f"{ROOT}/LABEL-Newsletter.xlsx")
if "Ranking_LLM_confronto" in wb.sheetnames: del wb["Ranking_LLM_confronto"]
ws = wb.create_sheet("Ranking_LLM_confronto")
hf = Font(bold=True, color="FFFFFF"); hd = PatternFill("solid", fgColor="7030A0")
green = PatternFill("solid", fgColor="E2EFDA"); red = PatternFill("solid", fgColor="FCE4D6")
yell = PatternFill("solid", fgColor="FFF2CC"); disf = PatternFill("solid", fgColor="FCE4EC")
ws["A1"] = "Confronto pertinenza LLM: gpt-5.4 vs gpt-5.6-sol (ordinato per score kNN)"
ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = f"5.4 Si={si4} | 5.6-sol Si={si6}. Rosa=disaccordo tra i due. Verde/Rosso=label nota."
ws["A2"].font = Font(italic=True, size=10, color="666666")
heads = ["Rank","Score kNN","Pert. 5.4","Pert. 5.6-sol","Accordo","Label nota","Titolo","Sommario","Fonte","Link"]
for c, h in enumerate(heads, 1):
    cell = ws.cell(row=4, column=c, value=h); cell.font = hf; cell.fill = hd
for rank, i in enumerate(order, 1):
    a = arts[i]; row = 4 + rank; t = a["title"]
    x4 = g54.get(t); x6 = p6map.get(t); L = lab.get(t)
    fonte, link = meta.get(t, ("", ""))
    ws.cell(row=row, column=1, value=rank)
    ws.cell(row=row, column=2, value=round(float(sc[i]), 4))
    c3 = ws.cell(row=row, column=3, value=("Si" if x4 == 1 else "No") if x4 is not None else "?")
    c4 = ws.cell(row=row, column=4, value=("Si" if x6 == 1 else "No") if x6 is not None else "?")
    acc = "OK" if x4 == x6 else "DIVERSO"
    c5 = ws.cell(row=row, column=5, value=acc)
    ws.cell(row=row, column=6, value=("Si" if L == 1 else "No") if L is not None else "")
    ws.cell(row=row, column=7, value=t)
    ws.cell(row=row, column=8, value=a["summary"][:300])
    ws.cell(row=row, column=9, value=fonte); ws.cell(row=row, column=10, value=link)
    if x4 == 1: c3.fill = yell
    if x6 == 1: c4.fill = yell
    if x4 != x6: c5.fill = disf
    if L == 0:
        for c in [6]: ws.cell(row=row, column=c).fill = red
    elif L == 1:
        for c in [6]: ws.cell(row=row, column=c).fill = green
    ws.cell(row=row, column=8).alignment = Alignment(wrap_text=True, vertical="top")
for col, w in {"A":6,"B":10,"C":10,"D":13,"E":10,"F":10,"G":55,"H":50,"I":18,"J":26}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
wb.save(f"{ROOT}/LABEL-Newsletter.xlsx")
print("Foglio 'Ranking_LLM_confronto' scritto.")
print("\nFINITO")
