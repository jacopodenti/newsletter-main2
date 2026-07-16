"""
Prototipo: classificatore di PERTINENZA OCX via Azure OpenAI (gpt-5.4).
Valuta sui 50 etichettati vs il giudizio della ricercatrice + misura token reali.
"""
import json, os, sys, time
from pathlib import Path
from dotenv import load_dotenv
from openai import OpenAI
import openpyxl

ROOT = Path("/Users/gabriele/Desktop/newsletter-main2")
load_dotenv(ROOT / ".env")
MODEL = "gpt-5.4"

client = OpenAI(api_key=os.environ["AZURE_OPENAI_API_KEY"], base_url=os.environ["AZURE_OPENAI_ENDPOINT"])

# --- descrizione OCX ---
oss = json.load(open(ROOT / "config" / "osservatori.json"))
ocx = next(o for o in oss if o["slug"] == "omnichannel-customer-experience")
ocx_desc = ocx.get("descrizione_enhanced") or ocx["descrizione"]

INSTR = f"""Sei un classificatore di pertinenza per l'Osservatorio "Omnichannel Customer Experience" (OCX) del Politecnico di Milano.

DESCRIZIONE OSSERVATORIO:
{ocx_desc}

COMPITO: dato un articolo (titolo + sommario), decidi se è PERTINENTE a questo osservatorio.
È pertinente SOLO se tratta in modo specifico di customer experience omnicanale, customer journey,
gestione della relazione col cliente su piu canali, contact center, CRM, self-service, loyalty, NPS.
NON e pertinente se parla genericamente di digitale, AI, cybersecurity, cloud, innovazione, startup
senza un focus concreto sull'esperienza del cliente.

Rispondi SOLO con un oggetto JSON: {{"rilevante": true/false, "confidence": 0.0-1.0, "motivo": "una frase in italiano"}}"""


def classify(title, summary):
    prompt = f"{INSTR}\n\nARTICOLO:\nTitolo: {title}\nSommario: {summary}\n\nJSON:"
    r = client.responses.create(model=MODEL, input=prompt)
    txt = r.output_text.strip()
    if txt.startswith("```"):
        txt = txt.split("```")[1].removeprefix("json").strip()
    return json.loads(txt), r.usage


# --- 50 etichettati (dal backup, cache intatta) ---
bk = openpyxl.load_workbook("LABEL-Newsletter.backup.xlsx", data_only=True)
rk = bk["Ranking"]
rows = []
for rr in range(3, rk.max_row + 1):
    h = rk.cell(row=rr, column=8).value
    if not h or str(h).strip() == "":
        continue
    rows.append((str(rk.cell(row=rr, column=3).value or ""),
                 str(rk.cell(row=rr, column=4).value or ""),
                 1 if str(h).strip().lower() == "si" else 0))
print(f"Articoli etichettati: {len(rows)}  (modello: {MODEL})")

tp = fp = fn = tn = 0
tot_in = tot_out = 0
errs = 0
preds = []
for i, (title, summary, y) in enumerate(rows, 1):
    try:
        res, usage = classify(title, summary)
        pred = 1 if res.get("rilevante") else 0
        tot_in += usage.input_tokens
        tot_out += usage.output_tokens
    except Exception as e:
        print(f"  [{i}] ERR {str(e)[:70]}"); errs += 1; continue
    preds.append((y, pred, res.get("confidence"), title))
    if pred and y: tp += 1
    elif pred and not y: fp += 1
    elif not pred and y: fn += 1
    else: tn += 1
    if i % 10 == 0:
        print(f"  {i}/{len(rows)}...")
    time.sleep(0.05)

n = tp + fp + fn + tn
prec = tp / (tp + fp) if (tp + fp) else 0
rec = tp / (tp + fn) if (tp + fn) else 0
f1 = 2 * prec * rec / (prec + rec) if (prec + rec) else 0
acc = (tp + tn) / n if n else 0

print("\n=== QUALITA vs ricercatrice ===")
print(f"  n valutati: {n}  (errori: {errs})")
print(f"  Accuracy : {acc:.0%}")
print(f"  Precision: {prec:.0%}   Recall: {rec:.0%}   F1: {f1:.2f}")
print(f"  Confusion: TP={tp} FP={fp} FN={fn} TN={tn}")

print("\n=== COSTO (token reali) ===")
print(f"  input tot: {tot_in}  output tot: {tot_out}  ({n} articoli)")
avg_in, avg_out = tot_in / max(n, 1), tot_out / max(n, 1)
print(f"  media/articolo: {avg_in:.0f} in + {avg_out:.0f} out")
wk_in, wk_out = avg_in * 1800, avg_out * 1800
print(f"  proiezione 1800/settimana: {wk_in/1e6:.2f}M input + {wk_out/1e6:.2f}M output token")

print("\n=== Falsi positivi (No detti Si) ===")
for y, p, conf, t in preds:
    if p == 1 and y == 0:
        print(f"  conf={conf}  {t[:65]}")
print("=== Falsi negativi (Si detti No) ===")
for y, p, conf, t in preds:
    if p == 0 and y == 1:
        print(f"  conf={conf}  {t[:65]}")
