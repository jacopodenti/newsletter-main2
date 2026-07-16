"""
Valutazione della classificazione CT LLM su Book-2.xlsx.

Produce:
  1. Accuratezza globale e per label (CT LLM vs Content Type ground truth)
  2. Matrice di confusione
  3. Correlazione Content Type × Rilevante OCX
  4. Lista casi di disaccordo

Output: stampa a terminale + report Word nella stessa cartella dell'Excel.

Uso:
    python src/evaluate_ct.py
    python src/evaluate_ct.py path/al/file.xlsx
"""
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).parent.parent
DEFAULT_EXCEL = ROOT / "Book-2.xlsx"
SHEET_NAME = "Foglio1"
HEADER_ROW = 1

LABELS = ["analysis", "case_study", "regulation", "research_report", "vendor_announcement"]
SHORT = {
    "analysis": "analysis",
    "case_study": "case_study",
    "regulation": "regulation",
    "research_report": "research_report",
    "vendor_announcement": "vendor_ann.",
}

# ─── colours ─────────────────────────────────────────────────────────────────
C_HEADER_BG  = "2E4057"   # dark blue
C_HEADER_FG  = "FFFFFF"
C_DIAG_BG    = "D6E4F0"   # light blue  (diagonal cells)
C_ALT_BG     = "F7F9FC"   # very light grey (alternating rows)
C_BORDER     = "B0BEC5"


# ─── data loading ─────────────────────────────────────────────────────────────

def load_data(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET_NAME]
    headers = {cell.value: cell.column - 1 for cell in ws[HEADER_ROW] if cell.value}

    required = {"Titolo", "Content Type", "CT LLM", "Rilevante"}
    missing = required - headers.keys()
    if missing:
        print(f"Errore: colonne mancanti: {missing}")
        sys.exit(1)

    rows = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        title = row[headers["Titolo"]]
        if not title or not str(title).strip():
            break
        rows.append(
            {
                "title": str(title).strip(),
                "gt": str(row[headers["Content Type"]] or "").strip(),
                "llm": str(row[headers["CT LLM"]] or "").strip(),
                "rilevante": str(row[headers["Rilevante"]] or "").strip(),
            }
        )
    return rows


# ─── compute ─────────────────────────────────────────────────────────────────

def compute_accuracy(rows: list[dict]) -> dict:
    total = len(rows)
    correct = sum(1 for r in rows if r["gt"] == r["llm"])
    per_label = {}
    for label in LABELS:
        gt_rows = [r for r in rows if r["gt"] == label]
        ok = sum(1 for r in gt_rows if r["llm"] == label)
        n = len(gt_rows)
        per_label[label] = {"n": n, "ok": ok, "acc": ok / n * 100 if n else 0.0}
    return {"total": total, "correct": correct, "pct": correct / total * 100 if total else 0.0, "per_label": per_label}


def compute_confusion(rows: list[dict]) -> dict[str, dict[str, int]]:
    matrix: dict[str, dict[str, int]] = {l: defaultdict(int) for l in LABELS}
    for r in rows:
        if r["gt"] in matrix:
            matrix[r["gt"]][r["llm"]] += 1
    return matrix


def compute_ocx(rows: list[dict]) -> dict[str, dict[str, int]]:
    counts: dict[str, dict[str, int]] = {l: {"Si": 0, "No": 0} for l in LABELS}
    for r in rows:
        label = r["gt"]
        if label in counts:
            key = r["rilevante"] if r["rilevante"] in ("Si", "No") else "No"
            counts[label][key] += 1
    return counts


def compute_disagreements(rows: list[dict]) -> list[dict]:
    return [r for r in rows if r["gt"] != r["llm"]]


# ─── terminal output ──────────────────────────────────────────────────────────

def _bar(n: int, total: int, width: int = 20) -> str:
    filled = round(width * n / total) if total else 0
    return "█" * filled + "░" * (width - filled)


def _section(title: str) -> None:
    print(f"\n{'═' * 62}")
    print(f"  {title}")
    print(f"{'═' * 62}")


def print_accuracy(acc: dict) -> None:
    _section("1. ACCURATEZZA CT LLM vs Content Type (ground truth)")
    print(f"\n  Totale articoli : {acc['total']}")
    print(f"  Match esatti    : {acc['correct']}  ({acc['pct']:.1f}%)")
    print(f"  Errori          : {acc['total'] - acc['correct']}  ({100 - acc['pct']:.1f}%)\n")
    print(f"  {'Label':<22} {'GT':>5} {'Corr':>5} {'Acc':>7}  Bar")
    print(f"  {'─' * 22} {'─' * 5} {'─' * 5} {'─' * 7}  {'─' * 20}")
    for label in LABELS:
        s = acc["per_label"][label]
        bar = _bar(s["ok"], s["n"]) if s["n"] else " " * 20
        print(f"  {label:<22} {s['n']:>5} {s['ok']:>5} {s['acc']:>6.1f}%  {bar}")


def print_confusion(matrix: dict) -> None:
    _section("2. MATRICE DI CONFUSIONE  (righe=GT, colonne=LLM)")
    col_w = 7
    short_labels = [SHORT[l] for l in LABELS]
    header = f"  {'GT \\ LLM':<22}" + "".join(f"{s[:col_w - 1]:>{col_w}}" for s in short_labels) + f"  {'TOT':>{col_w}}"
    print(f"\n{header}")
    print(f"  {'─' * 22}" + "─" * (col_w * len(LABELS)) + f"  {'─' * col_w}")
    for gt in LABELS:
        row_total = sum(matrix[gt].values())
        cells = ""
        for llm in LABELS:
            val = matrix[gt][llm]
            marker = f"[{val}]" if gt == llm else f" {val} "
            cells += f"{marker:>{col_w}}"
        print(f"  {gt:<22}{cells}  {row_total:>{col_w}}")
    print(f"  {'─' * 22}" + "─" * (col_w * len(LABELS)) + f"  {'─' * col_w}")
    totals = "".join(f"{sum(matrix[gt][llm] for gt in LABELS):>{col_w}}" for llm in LABELS)
    grand = sum(sum(matrix[gt].values()) for gt in LABELS)
    print(f"  {'TOT':<22}{totals}  {grand:>{col_w}}")
    print(f"\n  Valori sulla diagonale [n] = predizioni corrette.")


def print_ocx(counts: dict) -> None:
    _section("3. CORRELAZIONE Content Type × RILEVANTE OCX")
    print(f"\n  {'Label':<22} {'Si':>5} {'No':>5} {'Tot':>5}  {'% Si':>7}  Rilevanza")
    print(f"  {'─' * 22} {'─' * 5} {'─' * 5} {'─' * 5}  {'─' * 7}  {'─' * 20}")
    for label in LABELS:
        si = counts[label]["Si"]
        no = counts[label]["No"]
        tot = si + no
        pct = si / tot * 100 if tot else 0
        bar = _bar(si, tot) if tot else " " * 20
        print(f"  {label:<22} {si:>5} {no:>5} {tot:>5}  {pct:>6.1f}%  {bar}")


def print_disagreements(bad: list[dict]) -> None:
    _section("4. CASI DI DISACCORDO  (GT ≠ CT LLM)")
    if not bad:
        print("\n  Nessun disaccordo — classificazione perfetta.")
        return
    print(f"\n  {len(bad)} disaccordi\n")
    w = 52
    print(f"  #   {'Titolo':<{w}}  {'GT':<22}  CT LLM")
    print(f"  {'─' * 3} {'─' * w}  {'─' * 22}  {'─' * 22}")
    for i, r in enumerate(bad, 1):
        t = r["title"]
        t_short = t[: w - 1] + "…" if len(t) > w else t
        print(f"  {i:<3} {t_short:<{w}}  {r['gt']:<22}  {r['llm']}")


# ─── Word helpers ─────────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_border(cell, hex_color: str = C_BORDER) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), hex_color)
        tcBorders.append(border)
    tcPr.append(tcBorders)


def _cell_text(cell, text: str, bold: bool = False, color: str | None = None,
               align: str = "center", size: int = 10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        r, g, b = int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
        run.font.color.rgb = RGBColor(r, g, b)


def _header_row(table, values: list[str], col_widths: list[float] | None = None) -> None:
    row = table.rows[0]
    for i, (cell, val) in enumerate(zip(row.cells, values)):
        _set_cell_bg(cell, C_HEADER_BG)
        _set_cell_border(cell)
        _cell_text(cell, val, bold=True, color=C_HEADER_FG, align="center")
        if col_widths:
            cell.width = Cm(col_widths[i])


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.runs[0].font.color.rgb = RGBColor(0x2E, 0x40, 0x57)


# ─── Word sections ────────────────────────────────────────────────────────────

def _docx_accuracy(doc: Document, acc: dict) -> None:
    _add_heading(doc, "1. Accuratezza CT LLM vs Content Type", level=1)

    p = doc.add_paragraph()
    p.add_run(f"Totale articoli: {acc['total']}   |   "
              f"Match esatti: {acc['correct']} ({acc['pct']:.1f}%)   |   "
              f"Errori: {acc['total'] - acc['correct']} ({100 - acc['pct']:.1f}%)")
    p.runs[0].font.size = Pt(11)

    doc.add_paragraph()

    cols = ["Label", "Articoli GT", "Corretti", "Accuratezza"]
    widths = [5.0, 3.0, 3.0, 3.0]
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Table Grid"
    _header_row(table, cols, widths)

    for i, label in enumerate(LABELS):
        s = acc["per_label"][label]
        row = table.add_row()
        bg = C_ALT_BG if i % 2 else "FFFFFF"
        data = [label, str(s["n"]), str(s["ok"]), f"{s['acc']:.1f}%"]
        aligns = ["left", "center", "center", "center"]
        for cell, val, align in zip(row.cells, data, aligns):
            _set_cell_bg(cell, bg)
            _set_cell_border(cell)
            _cell_text(cell, val, align=align)


def _docx_confusion(doc: Document, matrix: dict) -> None:
    _add_heading(doc, "2. Matrice di confusione", level=1)
    doc.add_paragraph("Righe = Ground Truth (GT) · Colonne = CT LLM · "
                      "Celle evidenziate = predizioni corrette (diagonale)").runs[0].font.italic = True

    n = len(LABELS)
    table = doc.add_table(rows=1, cols=n + 2)
    table.style = "Table Grid"

    # header row: blank | labels... | TOT
    header_vals = ["GT \\ LLM"] + [SHORT[l] for l in LABELS] + ["TOT"]
    _header_row(table, header_vals)

    for i, gt in enumerate(LABELS):
        row = table.add_row()
        bg = C_ALT_BG if i % 2 else "FFFFFF"
        # first cell: GT label
        _set_cell_bg(row.cells[0], bg)
        _set_cell_border(row.cells[0])
        _cell_text(row.cells[0], SHORT[gt], bold=True, align="left")
        # counts
        row_total = sum(matrix[gt].values())
        for j, llm in enumerate(LABELS):
            val = matrix[gt][llm]
            cell = row.cells[j + 1]
            cell_bg = C_DIAG_BG if gt == llm else bg
            _set_cell_bg(cell, cell_bg)
            _set_cell_border(cell)
            _cell_text(cell, str(val), bold=(gt == llm), align="center")
        # TOT
        _set_cell_bg(row.cells[-1], bg)
        _set_cell_border(row.cells[-1])
        _cell_text(row.cells[-1], str(row_total), bold=True, align="center")

    # totals row
    tot_row = table.add_row()
    _set_cell_bg(tot_row.cells[0], C_HEADER_BG)
    _set_cell_border(tot_row.cells[0])
    _cell_text(tot_row.cells[0], "TOT", bold=True, color=C_HEADER_FG, align="left")
    grand = 0
    for j, llm in enumerate(LABELS):
        t = sum(matrix[gt][llm] for gt in LABELS)
        grand += t
        _set_cell_bg(tot_row.cells[j + 1], C_HEADER_BG)
        _set_cell_border(tot_row.cells[j + 1])
        _cell_text(tot_row.cells[j + 1], str(t), bold=True, color=C_HEADER_FG, align="center")
    _set_cell_bg(tot_row.cells[-1], C_HEADER_BG)
    _set_cell_border(tot_row.cells[-1])
    _cell_text(tot_row.cells[-1], str(grand), bold=True, color=C_HEADER_FG, align="center")


def _docx_ocx(doc: Document, counts: dict) -> None:
    _add_heading(doc, "3. Correlazione Content Type × Rilevante OCX", level=1)
    doc.add_paragraph("Per ogni label (ground truth): quanti articoli sono stati giudicati rilevanti per OCX.").runs[0].font.italic = True

    cols = ["Label", "Sì", "No", "Totale", "% Sì"]
    widths = [5.0, 2.0, 2.0, 2.5, 2.5]
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Table Grid"
    _header_row(table, cols, widths)

    for i, label in enumerate(LABELS):
        si = counts[label]["Si"]
        no = counts[label]["No"]
        tot = si + no
        pct = si / tot * 100 if tot else 0.0
        row = table.add_row()
        bg = C_ALT_BG if i % 2 else "FFFFFF"
        data = [label, str(si), str(no), str(tot), f"{pct:.1f}%"]
        aligns = ["left", "center", "center", "center", "center"]
        for cell, val, align in zip(row.cells, data, aligns):
            _set_cell_bg(cell, bg)
            _set_cell_border(cell)
            _cell_text(cell, val, align=align)


def _docx_disagreements(doc: Document, bad: list[dict], total: int) -> None:
    _add_heading(doc, "4. Casi di disaccordo (GT ≠ CT LLM)", level=1)

    if not bad:
        doc.add_paragraph("Nessun disaccordo — classificazione perfetta.")
        return

    doc.add_paragraph(f"{len(bad)} disaccordi su {total} articoli.").runs[0].font.italic = True

    cols = ["#", "Titolo", "Ground Truth", "CT LLM"]
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Table Grid"
    _header_row(table, cols)

    for i, r in enumerate(bad, 1):
        row = table.add_row()
        bg = C_ALT_BG if i % 2 else "FFFFFF"
        data = [str(i), r["title"], r["gt"], r["llm"]]
        aligns = ["center", "left", "left", "left"]
        for cell, val, align in zip(row.cells, data, aligns):
            _set_cell_bg(cell, bg)
            _set_cell_border(cell)
            _cell_text(cell, val, align=align)


# ─── Word export ──────────────────────────────────────────────────────────────

def write_docx(
    path: Path,
    rows: list[dict],
    acc: dict,
    matrix: dict,
    counts: dict,
    bad: list[dict],
) -> Path:
    doc = Document()

    # page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # title
    title = doc.add_heading("Report: valutazione classificazione CT LLM", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(0x2E, 0x40, 0x57)

    sub = doc.add_paragraph(
        f"File: {path.name}  ·  Articoli analizzati: {len(rows)}  ·  Data: {date.today().isoformat()}"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(10)
    sub.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    doc.add_paragraph()

    _docx_accuracy(doc, acc)
    doc.add_paragraph()
    _docx_confusion(doc, matrix)
    doc.add_paragraph()
    _docx_ocx(doc, counts)
    doc.add_paragraph()
    _docx_disagreements(doc, bad, total=len(rows))

    out = path.with_name(path.stem + "_ct_eval.docx")
    doc.save(out)
    return out


# ─── main ─────────────────────────────────────────────────────────────────────

def main() -> None:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_EXCEL
    if not path.exists():
        print(f"File non trovato: {path}")
        sys.exit(1)

    print(f"File: {path.name}  —  foglio: {SHEET_NAME}")
    rows = load_data(path)
    print(f"Righe caricate: {len(rows)}")

    acc    = compute_accuracy(rows)
    matrix = compute_confusion(rows)
    counts = compute_ocx(rows)
    bad    = compute_disagreements(rows)

    print_accuracy(acc)
    print_confusion(matrix)
    print_ocx(counts)
    print_disagreements(bad)

    out = write_docx(path, rows, acc, matrix, counts, bad)
    print(f"\nReport Word salvato: {out}")


if __name__ == "__main__":
    main()
