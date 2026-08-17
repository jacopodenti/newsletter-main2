"""
Backfill di selected_for_newsletter in store/articles.jsonl dalle 3 fonti editoriali
già mantenute a mano (le stesse che alimentano dataset.jsonl in <pipeline>/02_build_dataset.py)
— NESSUNA chiamata di rete/API, legge solo gli XLSX già sul disco.

Per OCX: solo le righe NON marcate come scartate nella colonna note (stessa lista di
parole di ocx_relevance/02_build_dataset.py::load_editorial() — "non va"/"fuori scope"/
"non bene"/"scartato" => non è una selezione). Per AI e Startup Thinking: tutte le righe,
perché le rispettive fonti editoriali non hanno mai un reject (documentato nei rispettivi
02_build_dataset.py: "tutti trattati come pertinenti").

Non riusa load_editorial() dei tre 02_build_dataset.py perché nessuno dei tre legge la
colonna Data (serve qui per first_seen/arco temporale, non serve a 02_build_dataset.py).

Uso: python backfill_selected_from_editorial.py
"""
import sys
from pathlib import Path
from datetime import datetime

import openpyxl

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "ocx_relevance"))
from common import log, canonicalize_url, load_tracking_params, upsert_articles_store

REJECT_WORDS = ["non va", "fuori scope", "non bene", "scartato"]


def to_iso(value) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str) and value.strip():
        return value.strip()[:10]
    return None


def read_ocx():
    """Tutti i fogli di OCX26_Newsletter.xlsx: colonne per nome header (posizione
    variabile da foglio a foglio, come già gestito in ocx_relevance/02_build_dataset.py)."""
    wb = openpyxl.load_workbook(ROOT / "OCX26_Newsletter.xlsx", data_only=True)
    out = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        hdr = {str(ws.cell(row=1, column=c).value).strip().lower(): c
               for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value}
        dc, lc = hdr.get("data"), hdr.get("link")
        nc = next((v for k, v in hdr.items() if k.startswith("note")), None)
        if not lc:
            continue
        for r in range(2, ws.max_row + 1):
            link = str(ws.cell(row=r, column=lc).value or "").strip()
            if not link:
                continue
            note = str(ws.cell(row=r, column=nc).value or "").strip().lower() if nc else ""
            if any(w in note for w in REJECT_WORDS):
                continue  # scartato in redazione: non è una selezione
            out.append((link, ws.cell(row=r, column=dc).value if dc else None))
    return out


def read_ai():
    """Unico foglio 'Worksheet' di ExportNewsLetters.xlsx: Titolo(1), Data(2), Link(4)."""
    wb = openpyxl.load_workbook(ROOT / "ai_relevance" / "ExportNewsLetters.xlsx", data_only=True)
    ws = wb["Worksheet"]
    out = []
    for r in range(2, ws.max_row + 1):
        link = str(ws.cell(row=r, column=4).value or "").strip()
        if not link:
            continue
        out.append((link, ws.cell(row=r, column=2).value))
    return out


def read_startup():
    """Fogli settimanali di SUT_Elenco news.xlsx (Data(1), LINK(3)) — salta i fogli di
    servizio in coda (Ranking_ST_nuovo, Analisi_Cluster, Costi), riconoscibili perché la
    loro intestazione in A1 non è 'Data'."""
    wb = openpyxl.load_workbook(ROOT / "startup_thinking_relevance" / "SUT_Elenco news.xlsx", data_only=True)
    out = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        if str(ws.cell(row=1, column=1).value or "").strip().lower() != "data":
            continue
        for r in range(2, ws.max_row + 1):
            link = str(ws.cell(row=r, column=3).value or "").strip()
            if not link:
                continue
            out.append((link, ws.cell(row=r, column=1).value))
    return out


def main():
    tp = load_tracking_params()
    sources = {"ocx": read_ocx(), "ai": read_ai(), "startup": read_startup()}

    all_dates = []
    total = 0
    log("Backfill selected_for_newsletter dalle fonti editoriali (zero chiamate API)")
    for osservatorio, rows in sources.items():
        records = []
        for link, raw_date in rows:
            iso = to_iso(raw_date)
            if iso:
                all_dates.append(iso)
            records.append({
                "url_canonico": canonicalize_url(link, tp),
                "osservatorio": osservatorio,
                "content_hash": None,
                "published_date": iso or "",
                "text_source": None,
                "last_rank": None,
                "first_seen": iso,  # data storica nota, non "oggi" (solo per righe NUOVE)
                "selected_for_newsletter": True,
            })
        stats = upsert_articles_store(records)
        log(f"  {osservatorio:8s}: {len(records):>4d} righe editoriali -> "
            f"{stats['nuove']} nuove, {stats['aggiornate']} già in store, totale store ora {stats['totale_dopo']}")
        total += len(records)

    log(f"\nTOTALE articoli importati (letti dalle fonti, prima di dedup su ripetizioni): {total}")
    if all_dates:
        log(f"arco temporale (published_date/Data disponibile su {len(all_dates)}/{total} righe): "
            f"{min(all_dates)} -> {max(all_dates)}")
    log("\nFATTO.")


if __name__ == "__main__":
    main()
